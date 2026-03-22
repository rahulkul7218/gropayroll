# Copyright (c) 2026, Antigravity and contributors
# For license information, please see license.txt
# Trigger reload

import frappe
from frappe import _
from frappe.utils import flt, getdate, format_date

def execute(filters=None):
	if not filters:
		filters = {}

	columns = get_columns()
	data = get_data(filters)
	
	# Calculate totals for the report UI if needed
	# Standard Frappe QueryReport will do this if add_total_row is 1 in JSON
	
	return columns, data

@frappe.whitelist()
def download_excel(filters):
	import openpyxl
	from openpyxl.styles import Font, PatternFill, Alignment
	from io import BytesIO
	
	if isinstance(filters, str):
		import json
		filters = json.loads(filters)
		
	columns, data = execute(filters)
	
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Monthly Salary Slip Register"
	
	# Header Style
	header_fill = PatternFill(start_color="D1D8E0", end_color="D1D8E0", fill_type="solid") # Light Greyish Blue
	header_font = Font(bold=True)
	
	# Add Filter Information at the top
	filters_to_show = [
		("Month", filters.get("month")),
		("From Date", filters.get("from_date")),
		("To Date", filters.get("to_date")),
	]
	
	for label, val in filters_to_show:
		if val:
			ws.append([label, val])
			last_row = ws.max_row
			ws.cell(row=last_row, column=1).font = Font(bold=True)
		
	ws.append([]) # Empty row for spacing
	
	# Add Columns
	column_labels = [c.get("label") for c in columns]
	ws.append(column_labels)
	
	header_row_idx = ws.max_row
	for i in range(1, len(column_labels) + 1):
		cell = ws.cell(row=header_row_idx, column=i)
		cell.fill = header_fill
		cell.font = header_font
		cell.alignment = Alignment(horizontal="center")
		
	# Add Data
	fieldnames = [c.get("fieldname") for c in columns]
	for row_data in data:
		row = [row_data.get(f) for f in fieldnames]
		ws.append(row)
		# Style Data Cells
		last_row_idx = ws.max_row
		for i in range(1, len(row) + 1):
			cell = ws.cell(row=last_row_idx, column=i)
			if isinstance(cell.value, (int, float)) and cell.value != 0:
				cell.number_format = '[>=100000]##\,##\,##0.00;##,##0.00' # Indian Format Lakhs
		
	# Add Total Row
	total_row = ["Total"] + [""] * (len(column_labels) - 1)
	totals = {}
	for f in fieldnames:
		if f in ["employee", "employee_name", "designation", "department", "unit_name", "remarks", "idx"]:
			continue
		totals[f] = sum(flt(row.get(f)) for row in data)
		
	# Build the total row list
	total_row = []
	for f in fieldnames:
		if f == "employee":
			total_row.append("Total")
		elif f in totals:
			total_row.append(totals[f])
		else:
			total_row.append("")
			
	ws.append(total_row)
	
	# Style Total Row (Last Row)
	total_row_idx = ws.max_row
	total_font = Font(bold=True)
	for i in range(1, len(fieldnames) + 1):
		cell = ws.cell(row=total_row_idx, column=i)
		cell.font = total_font
		if isinstance(cell.value, (int, float)):
			cell.number_format = '[>=100000]##\,##\,##0.00;##,##0.00' # Indian Format Lakhs

	# Adjust column widths
	for i, col in enumerate(ws.columns):
		max_length = 0
		column = col[0].column_letter
		for cell in col:
			try:
				if len(str(cell.value)) > max_length:
					max_length = len(str(cell.value))
			except:
				pass
		adjusted_width = (max_length + 2)
		ws.column_dimensions[column].width = min(adjusted_width, 30)

	# Save to stream
	output = BytesIO()
	wb.save(output)
	output.seek(0)
	
	frappe.response.filename = "Monthly_Salary_Slip_Register.xlsx"
	frappe.response.filecontent = output.getvalue()
	frappe.response.type = "binary"

def get_columns():
	return [
		# {"label": _("S.NO."), "fieldname": "idx", "fieldtype": "Int", "width": 50},
		{"label": _("EMPLOYEE CODE"), "fieldname": "employee", "fieldtype": "Data", "width": 120},
		{"label": _("EMPLOYEE NAME"), "fieldname": "employee_name", "fieldtype": "Data", "width": 150},
		{"label": _("DESIGNATION"), "fieldname": "designation", "fieldtype": "Data", "width": 120},
		{"label": _("DEPARTMENT"), "fieldname": "department", "fieldtype": "Data", "width": 120},
		{"label": _("UNIT NAME"), "fieldname": "unit_name", "fieldtype": "Data", "width": 120},
		{"label": _("STANDARD DAYS"), "fieldname": "standard_days", "fieldtype": "Float", "width": 100},
		{"label": _("GROSS SALARY"), "fieldname": "gross_salary", "fieldtype": "Currency", "width": 120},
		{"label": _("PRESENT DAYS"), "fieldname": "present_days", "fieldtype": "Float", "width": 100},
		{"label": _("WEEKLY OFF"), "fieldname": "weekly_off", "fieldtype": "Float", "width": 100},
		{"label": _("HOLIDAY"), "fieldname": "holiday", "fieldtype": "Float", "width": 100},
		{"label": _("LEAVES"), "fieldname": "leaves", "fieldtype": "Float", "width": 100},
		{"label": _("TOTAL DAYS"), "fieldname": "total_days", "fieldtype": "Float", "width": 100},
		{"label": _("BASIC"), "fieldname": "basic", "fieldtype": "Currency", "width": 120},
		{"label": _("HRA"), "fieldname": "hra", "fieldtype": "Currency", "width": 120},
		{"label": _("CONV. ALL."), "fieldname": "conv_all", "fieldtype": "Currency", "width": 120},
		{"label": _("SPL. ALL."), "fieldname": "spl_all", "fieldtype": "Currency", "width": 120},
		{"label": _("ALL.1"), "fieldname": "all_1", "fieldtype": "Currency", "width": 120},
		{"label": _("TOTAL RATE"), "fieldname": "total_rate", "fieldtype": "Currency", "width": 120},
		{"label": _("EPF"), "fieldname": "epf", "fieldtype": "Currency", "width": 120},
		{"label": _("ESI"), "fieldname": "esi", "fieldtype": "Currency", "width": 120},
		{"label": _("LOAN"), "fieldname": "loan", "fieldtype": "Currency", "width": 120},
		{"label": _("TDS"), "fieldname": "tds", "fieldtype": "Currency", "width": 120},
		{"label": _("PROFESSIONAL TAX"), "fieldname": "professional_tax", "fieldtype": "Currency", "width": 120},
		{"label": _("OTHER DEDUCTION"), "fieldname": "other_deduction", "fieldtype": "Currency", "width": 120},
		{"label": _("MISC"), "fieldname": "misc", "fieldtype": "Currency", "width": 120},
		{"label": _("TOTAL DEDUCTION"), "fieldname": "total_deduction", "fieldtype": "Currency", "width": 120},
		{"label": _("NET PAYABLE"), "fieldname": "net_payable", "fieldtype": "Currency", "width": 120},
		{"label": _("REMARKS"), "fieldname": "remarks", "fieldtype": "Data", "width": 150},
	]

def get_data(filters):
	data = []
	
	conditions = get_conditions(filters)
	
	# Fetch Salary Slips
	ss_list = frappe.db.get_all("Salary Slip",
		filters=conditions,
		fields=["name", "employee", "employee_name", "total_working_days", "gross_pay", "net_pay", "total_deduction", "branch", "department"]
	)
	
	# Pre-fetch Employee Details
	employee_map = {e.name: e for e in frappe.db.get_all("Employee", fields=["name", "holiday_list", "designation", "department", "all_1"])}
	
	# Fetch child table data
	for i, ss in enumerate(ss_list):
		ss_doc = frappe.get_doc("Salary Slip", ss.name)
		emp_info = employee_map.get(ss.employee, {})
		row = {
			"idx": i + 1,
			"employee": ss.employee,
			"employee_name": ss.employee_name,
			"designation": emp_info.get("designation"),
			"department": emp_info.get("department"),
			"unit_name": ss.branch or ss.department,
			"standard_days": flt(ss.total_working_days),
			"gross_salary": flt(ss.gross_pay),
			"total_deduction": flt(ss.total_deduction),
			"net_payable": flt(ss.net_pay),
			"present_days": flt(ss_doc.payment_days) if hasattr(ss_doc, "payment_days") else 0.0,
			"weekly_off": 0.0,
			"holiday": 0.0,
			"leaves": 0.0,
			"total_days": 0.0,
			"basic": 0.0,
			"hra": 0.0,
			"conv_all": 0.0,
			"spl_all": 0.0,
			"all_1": flt(emp_info.get("all_1")),
			"total_rate": flt(ss.gross_pay),
			"epf": 0.0,
			"esi": 0.0,
			"loan": 0.0,
			"tds": 0.0,
			"professional_tax": 0.0,
			"other_deduction": 0.0,
			"misc": 0.0,
			"remarks": getattr(ss_doc, "remarks", ""),
		}
		
		# Fetch Weekly Off and Holidays from Holiday List
		emp_info = employee_map.get(ss.employee, {})
		holiday_list = emp_info.get("holiday_list")
		if holiday_list:
			holidays = frappe.db.get_all("Holiday",
				filters=[
					["parent", "=", holiday_list],
					["holiday_date", ">=", ss_doc.start_date],
					["holiday_date", "<=", ss_doc.end_date]
				],
				fields=["weekly_off"]
			)
			
			for h in holidays:
				if h.weekly_off:
					row["weekly_off"] = flt(row["weekly_off"]) + 1
				else:
					row["holiday"] = flt(row["holiday"]) + 1
		# Fetch Approved Leaves from Leave Application
		leaves_list = frappe.db.get_all("Leave Application",
			filters=[
				["employee", "=", ss.employee],
				["status", "=", "Approved"],
				["from_date", "<=", ss_doc.end_date],
				["to_date", ">=", ss_doc.start_date]
			],
			fields=["from_date", "to_date", "total_leave_days", "half_day", "half_day_date"]
		)
		
		for l in leaves_list:
			# Calculate overlap period
			l_start = l.from_date if l.from_date > ss_doc.start_date else ss_doc.start_date
			l_end = l.to_date if l.to_date < ss_doc.end_date else ss_doc.end_date
			
			if l_start <= l_end:
				diff = frappe.utils.date_diff(l_end, l_start) + 1
				
				# Adjust for half-day if it falls within the relevant period
				if l.half_day and l.half_day_date and ss_doc.start_date <= l.half_day_date <= ss_doc.end_date:
					# If leave is only one day and it's half day, diff becomes 0.5
					# If it's a multi-day leave but one is half, diff becomes (n - 0.5)? 
					# Actually, 'total_leave_days' might be more accurate if fully contained.
					# But if leave spans multiple months, we need careful subtraction.
					# Standard Frappe approach: 
					# If application is fully in the period, just add total_leave_days
					if l.from_date >= ss_doc.start_date and l.to_date <= ss_doc.end_date:
						row["leaves"] = flt(row["leaves"]) + flt(l.total_leave_days)
					else:
						# Rough calculation for partial: subtract half day if the half day date is in range
						row["leaves"] = flt(row["leaves"]) + diff
						if l.half_day:
							row["leaves"] = flt(row["leaves"]) - 0.5
				else:
					row["leaves"] = flt(row["leaves"]) + diff
					
		row["total_days"] = flt(row["present_days"]) + flt(row["weekly_off"]) + flt(row["holiday"]) + flt(row["leaves"])
		
		# Fetch Earnings
		for e in ss_doc.earnings:
			comp = e.salary_component.lower()
			if "basic" in comp:
				row["basic"] = flt(e.amount)
			elif "hra" in comp or "house rent" in comp:
				row["hra"] = flt(e.amount)
			elif "conveyance" in comp or "conv" in comp:
				row["conv_all"] = flt(e.amount)
			elif "special" in comp or "spl" in comp:
				row["spl_all"] = flt(e.amount)
				
		# Fetch Deductions
		for d in ss_doc.deductions:
			comp = d.salary_component.lower()
			if "epf" in comp or "provident fund" in comp:
				row["epf"] = flt(d.amount)
			elif "esi" in comp:
				row["esi"] = flt(d.amount)
			elif "loan" in comp:
				row["loan"] = flt(d.amount)
			elif "tds" in comp or "income tax" in comp:
				row["tds"] = flt(d.amount)
			elif "professional tax" in comp or "prof tax" in comp:
				row["professional_tax"] = flt(d.amount)
			elif "other" in comp:
				row["other_deduction"] = flt(d.amount)
			elif "misc" in comp:
				row["misc"] = flt(d.amount)
				
		data.append(row)
		
	return data

def get_conditions(filters):
	conditions = {}
	conditions["docstatus"] = ["<", 2]
	
	if filters.get("from_date"):
		conditions["start_date"] = [">=", filters.get("from_date")]
		
	if filters.get("to_date"):
		conditions["end_date"] = ["<=", filters.get("to_date")]
		
	# Handle Month filter in JS if needed, but if it comes here:
	# Usually Month filter is used to set From/To Date.
	
	return conditions
